"""Compara la ESTRUCTURA de dos archivos de inventario (no los datos).

Uso:
    uv run python scripts/comparar_estructura.py <esperado.xlsx> <generado.xlsx>

Revisa hojas, encabezados, fila 1, formatos de la primera fila de datos,
formulas por columna, fila de subtotales, hoja COPIA y tablas dinamicas.
Lee la contrasena de EXCEL_PASSWORD (.env) si los archivos estan cifrados.
"""

from __future__ import annotations

import os
import re
import sys
from io import BytesIO
from pathlib import Path

_RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_RAIZ / "src"))

import openpyxl  # noqa: E402
from dotenv import load_dotenv  # noqa: E402

from insumos.adapters.excel.lector import ExcelReader  # noqa: E402

load_dotenv(_RAIZ / ".env")
HOJAS_DATOS = ("INVENTARIO", "INVENTARIO COPIA")


def _abrir(ruta: Path) -> openpyxl.Workbook:
    datos = ExcelReader(ruta=ruta, password=os.getenv("EXCEL_PASSWORD"))._bytes_planos()
    return openpyxl.load_workbook(BytesIO(datos))


def _forma(valor: object) -> str:
    """Formula sin numeros de fila (=G3*H3 -> =G#*H#), o el tipo del valor."""
    if isinstance(valor, str) and valor.startswith("="):
        return re.sub(r"(?<=[A-Z])\d+", "#", valor)
    return type(valor).__name__


def _resumen(wb: openpyxl.Workbook) -> dict[str, object]:
    r: dict[str, object] = {"hojas": wb.sheetnames}
    for nombre in HOJAS_DATOS:
        if nombre not in wb.sheetnames:
            r[nombre] = "FALTA"
            continue
        ws = wb[nombre]
        ancho = ws.max_column
        ultima = ws.max_row
        r[nombre] = {
            "columnas": ancho,
            "fila1": [_forma(ws.cell(1, c).value) for c in range(1, ancho + 1)],
            "encabezados": [ws.cell(2, c).value for c in range(1, ancho + 1)],
            "formato_fila3": [ws.cell(3, c).number_format for c in range(1, ancho + 1)],
            "tipos_fila3": [_forma(ws.cell(3, c).value) for c in range(1, ancho + 1)],
            "ultima_fila": [_forma(ws.cell(ultima, c).value) for c in range(1, ancho + 1)],
            "referencia_texto": ws.cell(3, 1).number_format == "@",
        }
    pivots = []
    for ws in wb.worksheets:
        for p in getattr(ws, "_pivots", []):
            src = p.cache.cacheSource.worksheetSource
            pivots.append((ws.title, p.name, getattr(src, "sheet", None)))
    r["tablas_dinamicas"] = sorted(pivots)
    return r


def main() -> int:
    esperado, generado = (Path(a) for a in sys.argv[1:3])
    a, b = _resumen(_abrir(esperado)), _resumen(_abrir(generado))
    diferencias = 0
    for clave in a:
        if isinstance(a[clave], dict) and isinstance(b.get(clave), dict):
            for sub, valor in a[clave].items():  # type: ignore[union-attr]
                otro = b[clave].get(sub)  # type: ignore[union-attr]
                if valor != otro:
                    diferencias += 1
                    print(f"[DIFERENTE] {clave}.{sub}\n   esperado: {valor}\n   generado: {otro}")
        elif a[clave] != b.get(clave):
            diferencias += 1
            print(f"[DIFERENTE] {clave}\n   esperado: {a[clave]}\n   generado: {b.get(clave)}")
    print("OK: misma estructura" if diferencias == 0 else f"{diferencias} diferencias de estructura")
    return 1 if diferencias else 0


if __name__ == "__main__":
    sys.exit(main())
