"""Reporte de eliminaciones de inventario, con el mismo formato del proceso anterior.

Nombre y hojas iguales al ``EliminacionTracker`` del legacy:

- ``REPORTE_ELIMINACIONES_<AAAAMMDD_HHMMSS>.xlsx`` en la carpeta de salida.
- FUENTES DE DATOS         archivo usado para cada insumo y su fecha (rojo si no es de hoy)
- TODAS LAS ELIMINACIONES  TIMESTAMP, PASO, FILA_EXCEL, REFERENCIA, NOMBRE, MARCA, LINEA, MOTIVO
- RESUMEN POR PASO         PASO, TOTAL ELIMINADAS, PORCENTAJE, EJEMPLOS
- una hoja por PASO

Se agregan dos hojas nuevas al final: REFERENCIAS NUEVAS (las que entraron
desde la base de datos) y RESUMEN REGLAS (antes/despues de cada regla).
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from insumos.domain.ports import EvidenciaEjecucion
from insumos.domain.rules.inventario import columnas as C

PREFIJO_REPORTE = "REPORTE_ELIMINACIONES"

# Nombre del PASO (como lo veia el negocio en el reporte anterior) por regla.
PASOS: dict[tuple[str, str], str] = {
    ("inventario_general", "inv.eliminar_referencias_duplicadas"): "PASO 1: REFERENCIA DUPLICADA",
    ("inventario_general", "inv.eliminar_ausentes_en_bd"): "PASO 2: NO LLEGA COMO INVENTARIO EN BASE DE DATOS",
    ("inventario_bd", "inv.filtrar_motivo_inventario"): "BD: EXCLUIDA POR MOTIVO",
    ("inventario_bd", "inv.eliminar_referencias_duplicadas"): "BD: REFERENCIA DUPLICADA",
}
COLUMNAS_ELIMINACIONES = ["TIMESTAMP", "PASO", "FILA_EXCEL", "REFERENCIA", "NOMBRE", "MARCA", "LINEA", "MOTIVO"]
_COLORES_PASO = ["FFF2CC", "E2EFDA", "DEEBF7", "FCE4D6", "EDEDED", "F4B084", "C5E0B4", "BDD7EE"]


def _texto(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    return str(v).strip()


def _primera(fila: pd.Series, *columnas: str) -> str:
    for col in columnas:
        if col in fila.index and _texto(fila[col]):
            return _texto(fila[col])
    return ""


def tabla_eliminaciones(evidencia: EvidenciaEjecucion, sello: str) -> pd.DataFrame:
    """Todas las eliminaciones de ambos pipelines en el formato del legacy."""
    registros: list[dict[str, Any]] = []
    for audit in (evidencia.audit_bd, evidencia.audit_inventario):
        tabla = audit.tabla_eliminaciones()
        for _, fila in tabla.iterrows():
            regla = _texto(fila.get("_regla"))
            motivo = _texto(fila.get("_motivo"))
            detalle = _texto(fila.get("_detalle"))
            fila_excel = fila.get("_FILA")
            registros.append({
                "TIMESTAMP": sello,
                "PASO": PASOS.get((audit.pipeline, regla), f"{audit.pipeline}: {regla}"),
                "FILA_EXCEL": int(fila_excel) if _texto(fila_excel) else None,
                "REFERENCIA": _texto(fila.get(C.REFERENCIA)),
                "NOMBRE": _primera(fila, C.NOMBRE_ODOO, C.NOMBRE_LISTA),
                "MARCA": _primera(fila, C.MARCA_COPIA, C.BD_MARCA, C.MARCA_SISTEMA),
                "LINEA": _primera(fila, C.LINEA_COPIA, C.BD_LINEA, C.LINEA_SISTEMA),
                "MOTIVO": f"{motivo}: {detalle}" if detalle else motivo,
            })
    df = pd.DataFrame(registros, columns=COLUMNAS_ELIMINACIONES)
    return df.sort_values(["PASO", "FILA_EXCEL"], kind="mergesort").reset_index(drop=True)


def tabla_resumen(eliminaciones: pd.DataFrame) -> pd.DataFrame:
    total = len(eliminaciones)
    filas = []
    for paso, grupo in eliminaciones.groupby("PASO", sort=False):
        ejemplos = ", ".join(grupo["REFERENCIA"].head(5).astype(str))
        if len(grupo) > 5:
            ejemplos += f" ... (+{len(grupo) - 5} mas)"
        filas.append({
            "PASO": paso,
            "TOTAL ELIMINADAS": len(grupo),
            "PORCENTAJE": f"{len(grupo) / total * 100:.1f}%" if total else "0.0%",
            "EJEMPLOS": ejemplos,
        })
    resumen = pd.DataFrame(filas, columns=["PASO", "TOTAL ELIMINADAS", "PORCENTAJE", "EJEMPLOS"])
    return resumen.sort_values("TOTAL ELIMINADAS", ascending=False)


@dataclass
class EscritorReporteXlsx:
    carpeta: Path

    def escribir(self, evidencia: EvidenciaEjecucion) -> Path:
        self.carpeta.mkdir(parents=True, exist_ok=True)
        ahora = datetime.now()
        ruta = self.carpeta / f"{PREFIJO_REPORTE}_{ahora:%Y%m%d_%H%M%S}.xlsx"

        eliminaciones = tabla_eliminaciones(evidencia, ahora.strftime("%Y-%m-%d %H:%M:%S"))
        datos = evidencia.datos
        nuevas = datos[datos["_NUEVA"].astype(bool)] if "_NUEVA" in datos.columns else datos.iloc[0:0]
        nuevas = nuevas[[c for c in C.COLUMNAS_PLANTILLA if c in nuevas.columns and c not in C.COLUMNAS_FORMULA]]
        reglas = pd.concat(
            [
                evidencia.audit_bd.tabla_pasos().assign(PIPELINE=evidencia.audit_bd.pipeline),
                evidencia.audit_inventario.tabla_pasos().assign(PIPELINE=evidencia.audit_inventario.pipeline),
            ],
            ignore_index=True,
        )

        with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
            pd.DataFrame(
                list(evidencia.fuentes), columns=["DATO", "ARCHIVO", "FECHA MODIFICACION", "ADVERTENCIA"]
            ).to_excel(xw, sheet_name="FUENTES DE DATOS", index=False)
            eliminaciones.to_excel(xw, sheet_name="TODAS LAS ELIMINACIONES", index=False)
            tabla_resumen(eliminaciones).to_excel(xw, sheet_name="RESUMEN POR PASO", index=False)
            for paso in sorted(eliminaciones["PASO"].unique()):
                nombre = paso.replace(":", "").replace("/", "-")[:31]
                eliminaciones[eliminaciones["PASO"] == paso].to_excel(xw, sheet_name=nombre, index=False)
            nuevas.to_excel(xw, sheet_name="REFERENCIAS NUEVAS", index=False)
            reglas.to_excel(xw, sheet_name="RESUMEN REGLAS", index=False)

        _dar_formato(ruta)
        return ruta


def _dar_formato(ruta: Path) -> None:
    """Mismo formato del reporte anterior: anchos, encabezado azul, colores por PASO."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = load_workbook(ruta)
    for ws in wb.worksheets:
        for columna in ws.columns:
            largo = max((len(str(c.value)) for c in columna if c.value is not None), default=0)
            ws.column_dimensions[columna[0].column_letter].width = min(largo + 2, 50)
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            celda.alignment = Alignment(horizontal="center", vertical="center")
        encabezados = {c.value: c.column for c in ws[1]}
        if ws.title == "FUENTES DE DATOS" and "ADVERTENCIA" in encabezados:
            col = encabezados["ADVERTENCIA"]
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, col).value:
                    for c in range(1, ws.max_column + 1):
                        ws.cell(r, c).fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
                    ws.cell(r, col).font = Font(bold=True, color="C00000")
        if ws.title == "TODAS LAS ELIMINACIONES" and "PASO" in encabezados:
            col = encabezados["PASO"]
            colores: dict[str, str] = {}
            for r in range(2, ws.max_row + 1):
                paso = ws.cell(r, col).value
                if not paso:
                    continue
                color = colores.setdefault(paso, _COLORES_PASO[len(colores) % len(_COLORES_PASO)])
                relleno = PatternFill(start_color=color, end_color=color, fill_type="solid")
                for c in range(1, ws.max_column + 1):
                    ws.cell(r, c).fill = relleno
    wb.save(ruta)
