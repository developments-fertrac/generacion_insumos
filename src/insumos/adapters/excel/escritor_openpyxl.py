"""Escritor de inventario sobre la plantilla real, sin Excel (openpyxl).

Produce un archivo con la **misma estructura** que la salida de produccion:
todas las hojas de la plantilla, fila 1 con sus formulas, encabezados en la
fila 2, formatos de la primera fila de datos, TOTAL INV y Dif como formulas,
fila de subtotales al final, INVENTARIO COPIA como respaldo en valores y las
tablas dinamicas apuntando al nuevo rango (se refrescan al abrir).

Se usa en ``--dry-run`` (revision antes de produccion) y es el candidato a
reemplazar al escritor COM (plan, fase 6): no necesita Excel instalado.

Limitacion: openpyxl no calcula formulas. Las celdas con formula quedan sin
valor en cache hasta que Excel abre y recalcula el archivo; la hoja COPIA si
lleva los valores (se calculan aqui).
"""

from __future__ import annotations

import logging
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd

from insumos.adapters.excel.com_inventario import (
    _cifrar,
    ajustar_formula_subtotal,
    construir_matriz,
    letra_columna,
    mapear_columnas,
    nombre_salida,
)
from insumos.adapters.excel.escritor_simple import titulo_existencia
from insumos.adapters.excel.lector import ExcelReader
from insumos.domain.rules.inventario import columnas as C

log = logging.getLogger("actualizacion_inventario")


@dataclass
class EscritorInventarioOpenpyxl:
    """Implementa ``EscritorInventario`` editando la plantilla con openpyxl."""

    plantilla: Path
    carpeta_salida: Path
    password: str | None
    prefijo_salida: str = "DRYRUN INVENTARIO GENERAL"
    cifrar_salida: bool = True
    hoja: str = "INVENTARIO"
    hoja_copia: str = "INVENTARIO COPIA"
    fila_encabezado: int = 2

    def escribir(self, datos: pd.DataFrame, *, hoy: date) -> Path:
        import openpyxl

        contenido = BytesIO(ExcelReader(ruta=self.plantilla, password=self.password)._bytes_planos())
        wb = openpyxl.load_workbook(contenido, keep_links=True)
        if self.hoja not in wb.sheetnames:
            raise ValueError(f"La plantilla no tiene hoja {self.hoja}: {wb.sheetnames}")
        ws = wb[self.hoja]

        ultima, ancho = self._proyectar(ws, datos, hoy)
        self._clonar_copia(wb, ws, datos, ultima, ancho)
        self._actualizar_tablas_dinamicas(wb, ultima, ancho)

        self.carpeta_salida.mkdir(parents=True, exist_ok=True)
        destino = self.carpeta_salida / nombre_salida(self.prefijo_salida, datetime.now())
        wb.save(destino)
        if self.cifrar_salida and self.password:
            _cifrar(destino, self.password)
        log.info("Inventario escrito (openpyxl): %s (%d referencias)", destino, len(datos))
        return destino

    # ------------------------------------------------------------------

    def _proyectar(self, ws: Any, datos: pd.DataFrame, hoy: date) -> tuple[int, int]:
        h = self.fila_encabezado
        ancho = ws.max_column
        columnas = mapear_columnas([ws.cell(h, c).value for c in range(1, ancho + 1)])
        ultima_anterior = ws.max_row

        # Estilos de referencia: primera fila de datos y fila de subtotales de ayer.
        estilo_dato = [copy(ws.cell(h + 1, c)._style) for c in range(1, ancho + 1)]
        estilo_subtotal = [copy(ws.cell(ultima_anterior, c)._style) for c in range(1, ancho + 1)]
        alto_dato = ws.row_dimensions[h + 1].height

        ws.cell(h, columnas[C.EXISTENCIA]).value = titulo_existencia(hoy)
        if ultima_anterior > h:
            ws.delete_rows(h + 1, ultima_anterior - h)

        primera = h + 1
        ultima = h + len(datos)
        for i, fila in enumerate(construir_matriz(datos, columnas, ancho, primera)):
            r = primera + i
            for c, valor in enumerate(fila, start=1):
                celda = ws.cell(r, c)
                celda.value = valor
                celda._style = copy(estilo_dato[c - 1])
            if alto_dato is not None:
                ws.row_dimensions[r].height = alto_dato

        fila_subtotal = ultima + 1
        for c in range(1, ancho + 1):
            ws.cell(fila_subtotal, c)._style = copy(estilo_subtotal[c - 1])
        for nombre in (C.EXISTENCIA, C.TOTAL_INV):
            letra = letra_columna(columnas[nombre])
            ws.cell(fila_subtotal, columnas[nombre]).value = f"=SUBTOTAL(109,{letra}{primera}:{letra}{ultima})"

        for fila in range(1, h):
            for c in range(1, ancho + 1):
                valor = ws.cell(fila, c).value
                if isinstance(valor, str) and "SUBTOTAL(" in valor.upper():
                    ws.cell(fila, c).value = ajustar_formula_subtotal(valor, primera, ultima)

        self._columnas = columnas
        return ultima, ancho

    def _clonar_copia(self, wb: Any, ws: Any, datos: pd.DataFrame, ultima: int, ancho: int) -> None:
        """INVENTARIO COPIA: misma forma que INVENTARIO pero en valores calculados."""
        h = self.fila_encabezado
        anterior: dict[tuple[int, int], Any] = {}
        if self.hoja_copia in wb.sheetnames:
            vieja = wb[self.hoja_copia]
            anterior = {(f, c): vieja.cell(f, c).value for f in range(1, h) for c in range(1, ancho + 1)}
            indice = wb.sheetnames.index(self.hoja_copia)
            wb.remove(vieja)
        else:
            indice = wb.sheetnames.index(self.hoja)
        copia = wb.create_sheet(self.hoja_copia, index=indice)

        for letra, dim in ws.column_dimensions.items():
            copia.column_dimensions[letra].width = dim.width

        existencia = pd.to_numeric(datos[C.EXISTENCIA], errors="coerce").fillna(0)
        total = pd.to_numeric(datos[C.TOTAL_INV], errors="coerce").fillna(0)
        sumas = {
            letra_columna(self._columnas[C.EXISTENCIA]): float(existencia.sum()),
            letra_columna(self._columnas[C.TOTAL_INV]): float(total.sum()),
        }

        def _valor_calculado(f: int, c: int, valor: Any) -> Any:
            if isinstance(valor, str) and valor.startswith("="):
                if "SUBTOTAL(" in valor.upper():
                    return sumas.get(letra_columna(c))
                return anterior.get((f, c), 0)
            return valor

        # Filas de cabecera (1..h)
        for f in range(1, h + 1):
            for c in range(1, ancho + 1):
                origen = ws.cell(f, c)
                destino = copia.cell(f, c, _valor_calculado(f, c, origen.value))
                destino._style = copy(origen._style)

        # Datos en valores
        primera = h + 1
        filas = construir_matriz(datos, self._columnas, ancho, primera, formulas=False)
        for i, fila in enumerate(filas):
            r = primera + i
            for c, valor in enumerate(fila, start=1):
                celda = copia.cell(r, c, valor)
                celda._style = copy(ws.cell(r, c)._style)

        # Subtotales en valores
        fila_subtotal = ultima + 1
        for c in range(1, ancho + 1):
            origen = ws.cell(fila_subtotal, c)
            celda = copia.cell(fila_subtotal, c, _valor_calculado(fila_subtotal, c, origen.value))
            celda._style = copy(origen._style)

    def _actualizar_tablas_dinamicas(self, wb: Any, ultima: int, ancho: int) -> None:
        rango = f"A{self.fila_encabezado}:{letra_columna(ancho)}{ultima}"
        for hoja in wb.worksheets:
            for pivot in getattr(hoja, "_pivots", []):
                try:
                    fuente = pivot.cache.cacheSource.worksheetSource
                except AttributeError:
                    continue
                if fuente is None or str(fuente.sheet).strip().upper() != self.hoja.upper():
                    continue
                fuente.ref = rango
                pivot.cache.refreshOnLoad = True
                log.info("Tabla dinamica '%s' (%s) -> %s!%s", pivot.name, hoja.title, self.hoja, rango)
