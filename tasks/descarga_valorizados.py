"""
Tarea: Descarga de Valorizados por Almacen desde ERP Fertrac
Migrado desde Descargar Valorizados/descargar_valorizados.py
"""

from __future__ import annotations

import glob
import os
import shutil
import time
from datetime import datetime

from selenium.webdriver.common.by import By
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select, WebDriverWait

from config.erp_selectors import (
    AlertaERP,
    Menu,
    Modal,
    Navegacion,
    Valorizado,
)
from config.settings import Settings
from core.browser import create_driver, do_login, wait_for_download
from core.email_notifier import EmailNotifier
from core.erp_navigation import (
    click_with_fallback,
    close_modal,
    descartar_alert_js,
    detect_and_accept_empty_alert,
    find_all_visible,
    find_visible_element,
    open_menu_by_selectors,
    wait_for_erp_section,
)
from tasks.base_task import BaseTask


class DescargaValorizados(BaseTask):
    name = "descarga_valorizados"

    ALMACENES_CONFIG = [
        {
            "tipo_consulta": None,
            "ubicacion": None,
            "nombre_archivo": "VALORIZADO GENERAL.xlsx",
        },
        {
            "tipo_consulta": "Ubicación",
            "ubicacion": "3/Aforo Impo",
            "nombre_archivo": "VALORIZADO TOBERIN.xlsx",
        },
        {
            "tipo_consulta": "Ubicación",
            "ubicacion": "4/Faltantes",
            "nombre_archivo": "VALORIZADO FALTANTES.xlsx",
        },
        {
            "tipo_consulta": "Ubicación",
            "ubicacion": "7/Faltantes_Impo",
            "nombre_archivo": "VALORIZADO FALTANTES IMPO.xlsx",
        },
    ]

    def __init__(self, settings: Settings):
        super().__init__(settings)
        self._driver = None

    def setup(self) -> None:
        super().setup()
        self.notifier = EmailNotifier(self.settings.smtp, self.name)
        self.paths = self.settings.paths
        self.erp = self.settings.erp
        self.browser_cfg = self.settings.browser

    def execute(self) -> None:
        tiempo_inicio_total = time.time()
        archivos_descargados = []

        carpeta_destino = str(self.paths.valorizados)
        os.makedirs(carpeta_destino, exist_ok=True)
        self.log.info("Carpeta de destino: %s", carpeta_destino)

        self._driver = create_driver(carpeta_destino, self.browser_cfg)
        driver = self._driver

        if not do_login(driver, self.erp, self.browser_cfg):
            raise RuntimeError("Fallo en el login")

        self.navegar_a_inventario(driver)

        if not self.abrir_menu_informes(driver):
            raise RuntimeError("Fallo al abrir menu Informes")

        if not self.seleccionar_valorizado(driver):
            raise RuntimeError("Fallo al seleccionar Valorizado")

        self.log.info("=" * 70)
        self.log.info("DESCARGANDO INFORMES")
        self.log.info("=" * 70)

        for i, config in enumerate(self.ALMACENES_CONFIG, 1):
            self.log.info(
                "Procesando %d/%d: %s", i, len(self.ALMACENES_CONFIG), config["nombre_archivo"]
            )

            if config["tipo_consulta"] is not None:
                if not self.seleccionar_tipo_consulta(driver, config["tipo_consulta"]):
                    self.log.warning(
                        "No se pudo seleccionar tipo consulta '%s', continuando...",
                        config["tipo_consulta"],
                    )
                    continue
            else:
                self.log.info("Usando tipo de consulta predeterminado (Compañía)")
                time.sleep(2)

            if config["tipo_consulta"] == "Ubicación" and config["ubicacion"]:
                if not self.seleccionar_ubicacion_dropdown(driver, config["ubicacion"]):
                    self.log.warning(
                        "No se pudo seleccionar ubicacion '%s', continuando...",
                        config["ubicacion"],
                    )
                    continue

            archivos_antes = {}
            for ext in ("*.xlsx", "*.xls"):
                for f in glob.glob(os.path.join(carpeta_destino, ext)):
                    if not os.path.basename(f).startswith("~$"):
                        archivos_antes[f] = os.path.getmtime(f)
            self.log.info("Snapshot: %d archivos preexistentes", len(archivos_antes))

            if not self.generar_xlsx(driver):
                self.log.warning(
                    "No se pudo generar XLSX para '%s', continuando...", config["nombre_archivo"]
                )
                continue

            if detect_and_accept_empty_alert(driver, timeout=5):
                archivo_vacio = self._crear_valorizado_vacio(
                    carpeta_destino, config["nombre_archivo"]
                )
                if archivo_vacio:
                    archivos_descargados.append(archivo_vacio)
                    self.log.info(
                        "%s generado como vacío (bodega sin unidades)", config["nombre_archivo"]
                    )
                else:
                    self.log.warning(
                        "No se pudo crear archivo vacío para '%s'", config["nombre_archivo"]
                    )
            else:
                archivo_descargado = wait_for_download(
                    carpeta_destino,
                    timeout=self.browser_cfg.timeout_descarga,
                    snapshot=archivos_antes,
                )

                if archivo_descargado:
                    time.sleep(2)
                    archivo_final = self.renombrar_y_mover_archivo(
                        archivo_descargado,
                        config["nombre_archivo"],
                        carpeta_destino,
                    )
                    archivos_descargados.append(archivo_final)
                    self.log.info("%s descargado y renombrado", config["nombre_archivo"])
                else:
                    self.log.warning(
                        "Timeout esperando descarga de '%s'", config["nombre_archivo"]
                    )

            if i < len(self.ALMACENES_CONFIG):
                self.log.info("Preparando para siguiente descarga...")
                close_modal(driver)
                time.sleep(1)
                if not self.abrir_menu_informes(driver):
                    self.log.warning("No se pudo reabrir menu Informes")
                    break
                if not self.seleccionar_valorizado(driver):
                    self.log.warning("No se pudo reabrir Valorizado")
                    break

        tiempo_total_seg = int(time.time() - tiempo_inicio_total)
        tiempo_total_texto = (
            f"{tiempo_total_seg // 60} minutos {tiempo_total_seg % 60} segundos"
        )

        self.log.info("PROCESO COMPLETADO")
        self.log.info(
            "Archivos descargados: %d/%d", len(archivos_descargados), len(self.ALMACENES_CONFIG)
        )
        for archivo in archivos_descargados:
            self.log.info("  - %s", os.path.basename(archivo))
        self.log.info("Ubicacion: %s", carpeta_destino)
        self.log.info("Tiempo total: %s", tiempo_total_texto)

        nombres_descargados = [os.path.basename(a) for a in archivos_descargados]
        archivos_no_descargados = [
            c["nombre_archivo"]
            for c in self.ALMACENES_CONFIG
            if c["nombre_archivo"] not in nombres_descargados
        ]

        self._enviar_email_exito(
            archivos_descargados=archivos_descargados,
            archivos_fallidos=archivos_no_descargados if archivos_no_descargados else None,
            tiempo_total=tiempo_total_texto,
        )

    def teardown(self) -> None:
        if self._driver:
            try:
                self._driver.quit()
                self.log.info("Navegador cerrado")
            except Exception:
                pass
            self._driver = None
        self.log.info("Finalizado: %s", datetime.now().strftime("%d/%m/%Y %H:%M:%S"))

    # ── Email helpers ──────────────────────────────────────────────────

    def _enviar_email_exito(
        self,
        archivos_descargados: list[str],
        archivos_fallidos: list[str] | None = None,
        tiempo_total: str = "",
    ) -> None:
        fecha = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        archivos_html = "".join(
            f"<li>✅ {os.path.basename(a)}</li>" for a in archivos_descargados
        )
        fallidos_html = ""
        if archivos_fallidos:
            fallidos_html = (
                "<h3>⚠️ No descargados:</h3><ul>"
                + "".join(f"<li>❌ {a}</li>" for a in archivos_fallidos)
                + "</ul>"
            )

        html = f"""
        <html><body>
        <h2>✅ Descarga de Valorizados completada</h2>
        <p><b>Fecha:</b> {fecha}</p>
        <p><b>Tiempo total:</b> {tiempo_total}</p>
        <h3>Archivos descargados:</h3><ul>{archivos_html}</ul>
        {fallidos_html}
        <hr><p><small>Mensaje automático — {self.name}</small></p>
        </body></html>
        """
        self.notifier.send(
            subject=f"✅ Valorizados FERTRAC - Descarga completada {fecha}",
            html_body=html,
        )

    # ── Navegación Selenium ────────────────────────────────────────────

    def navegar_a_inventario(self, driver) -> None:
        self.log.info("Navegando a Inventario...")
        driver.get(self.erp.url_inventario)

        try:
            WebDriverWait(driver, 30).until(
                EC.presence_of_element_located(("xpath", Navegacion.INVENTARIO))
            )
            self.log.info("En la seccion de Inventario")

            self.log.info("Esperando carga completa de Inventario...")
            WebDriverWait(driver, 20).until(
                EC.presence_of_element_located(
                    ("xpath", "//*[contains(text(),'Informes') or contains(text(),'informes')]")
                )
            )
            self.log.info("Menu visible, Inventario cargado")
        except Exception:
            self.log.warning("Timeout esperando menu, continuando de todas formas...")

    def abrir_menu_informes(self, driver) -> bool:
        self.log.info("Abriendo menu 'Informes'...")

        for selector in Menu.INFORMES:
            try:
                self.log.info("Buscando 'Informes' con: %s...", selector[:70])

                elem = WebDriverWait(driver, 15).until(
                    EC.element_to_be_clickable(("xpath", selector))
                )

                if elem.is_displayed():
                    driver.execute_script("arguments[0].scrollIntoView(true);", elem)
                    time.sleep(0.3)
                    click_with_fallback(driver, elem, "Informes")

                    self.log.info("Menu 'Informes' abierto")

                    WebDriverWait(driver, 10).until(
                        EC.presence_of_element_located(
                            ("xpath", "//*[contains(text(),'Valorizado')]")
                        )
                    )
                    return True
            except Exception:
                continue

        self.log.error("No se encontro el menu 'Informes'")
        return False

    def seleccionar_valorizado(self, driver, max_intentos: int = 3) -> bool:
        self.log.info("Seleccionando 'Valorizado'...")

        for intento in range(1, max_intentos + 1):
            self.log.info("Intento %d/%d...", intento, max_intentos)

            for selector in Valorizado.OPCION:
                try:
                    self.log.info("Buscando 'Valorizado' con: %s...", selector[:60])

                    elem = WebDriverWait(driver, 10).until(
                        EC.element_to_be_clickable(("xpath", selector))
                    )

                    if elem.is_displayed():
                        click_with_fallback(driver, elem, "Valorizado")

                        self.log.info("'Valorizado' seleccionado - Esperando modal...")

                        WebDriverWait(driver, 15).until(
                            EC.presence_of_element_located(("xpath", Modal.CONTENEDOR))
                        )
                        self.log.info("Modal abierto")
                        return True
                except Exception:
                    continue

            if intento < max_intentos:
                self.log.warning(
                    "No se encontró 'Valorizado', esperando 5s antes de reintentar..."
                )
                time.sleep(5)

        self.log.error(
            "No se encontró la opción 'Valorizado' después de todos los intentos"
        )
        return False

    def seleccionar_tipo_consulta(self, driver, tipo_consulta: str) -> bool:
        self.log.info("Seleccionando '%s' en 'Tipo de consulta'...", tipo_consulta)

        try:
            WebDriverWait(driver, 10).until(
                EC.presence_of_element_located(("xpath", "//div[contains(@class,'modal')]//select"))
            )

            for selector in Valorizado.TIPO_CONSULTA_DROPDOWN:
                elementos = driver.find_elements(By.XPATH, selector)
                for elem in elementos:
                    if elem.is_displayed():
                        try:
                            select = Select(elem)
                            opciones_texto = [opt.text for opt in select.options]
                            if any(
                                tipo_consulta.lower() in opt.lower()
                                for opt in opciones_texto
                            ):
                                self.log.info("Dropdown 'Tipo de consulta' encontrado")
                                for opcion in select.options:
                                    if tipo_consulta.lower() in opcion.text.lower():
                                        self.log.info(
                                            "Seleccionando opcion: '%s'", opcion.text
                                        )
                                        select.select_by_visible_text(opcion.text)
                                        time.sleep(1)
                                        self.log.info("'%s' seleccionado", tipo_consulta)
                                        return True
                        except Exception:
                            continue

            raise RuntimeError(
                f"No se encontro el dropdown o la opcion '{tipo_consulta}'"
            )

        except Exception as e:
            self.log.error("Error seleccionando tipo consulta: %s", e)
            return False

    def seleccionar_ubicacion_dropdown(self, driver, nombre_ubicacion: str) -> bool:
        self.log.info("Seleccionando ubicacion '%s'...", nombre_ubicacion)

        tiempo_inicio = time.time()
        timeout_total = 30

        try:
            time.sleep(2)

            selects_en_modal = driver.find_elements(
                By.XPATH, "//div[contains(@class, 'modal')]//select"
            )
            selects_visibles = [s for s in selects_en_modal if s.is_displayed()]
            self.log.info("%d SELECT(s) visible(s)", len(selects_visibles))

            if len(selects_visibles) == 1:
                self.log.info(
                    "El campo de Ubicación NO es un SELECT, buscando INPUT..."
                )
                try:
                    descartar_alert_js(driver)
                    # Localizar el input con busquedas acotadas (sin waits largos
                    # que puedan parecer/convertirse en un cuelgue).
                    campo_ubicacion = None
                    if time.time() - tiempo_inicio < timeout_total:
                        try:
                            inputs_especificos = [
                                i for i in driver.find_elements(
                                    By.XPATH, Valorizado.UBICACION_INPUT_NOMBRE
                                ) if i.is_displayed()
                            ]
                            if inputs_especificos:
                                campo_ubicacion = inputs_especificos[0]
                                self.log.info("Usando input especifico location_id")
                        except Exception:
                            campo_ubicacion = None

                    if campo_ubicacion is None and time.time() - tiempo_inicio < timeout_total:
                        try:
                            inputs = driver.find_elements(By.XPATH, Valorizado.UBICACION_INPUT)
                            inputs_visibles = [inp for inp in inputs if inp.is_displayed()]
                            self.log.info("%d INPUT(s) visible(s)", len(inputs_visibles))
                            campo_ubicacion = (
                                inputs_visibles[-1]
                                if len(inputs_visibles) >= 2
                                else (inputs_visibles[0] if inputs_visibles else None)
                            )
                        except Exception:
                            campo_ubicacion = None

                    if campo_ubicacion and time.time() - tiempo_inicio < timeout_total:
                        driver.execute_script(
                            "arguments[0].scrollIntoView({block: 'center'});",
                            campo_ubicacion,
                        )
                        driver.execute_script(
                            "arguments[0].value = '';", campo_ubicacion
                        )
                        time.sleep(0.3)
                        click_with_fallback(driver, campo_ubicacion, "campo ubicacion")
                        time.sleep(0.5)

                        descartar_alert_js(driver)
                        campo_ubicacion.send_keys(nombre_ubicacion)

                        opciones_encontradas = []
                        for selector in Valorizado.AUTOCOMPLETE_OPCIONES:
                            if time.time() - tiempo_inicio > timeout_total:
                                break
                            try:
                                descartar_alert_js(driver)
                                WebDriverWait(driver, 3).until(
                                    EC.presence_of_element_located(("xpath", selector))
                                )
                                opciones = driver.find_elements(By.XPATH, selector)
                                opciones_visibles = [
                                    opt for opt in opciones if opt.is_displayed()
                                ]
                                if opciones_visibles:
                                    opciones_encontradas = opciones_visibles
                                    break
                            except Exception:
                                continue

                        opcion_buscar_mas = None
                        for opcion in opciones_encontradas:
                            try:
                                texto = self._texto_opcion(opcion)
                                if not texto:
                                    continue
                                if texto == nombre_ubicacion:
                                    opcion.click()
                                    time.sleep(0.5)
                                    self.log.info(
                                        "Ubicación '%s' seleccionada",
                                        nombre_ubicacion,
                                    )
                                    return True
                                if "Buscar" in texto and "más" in texto:
                                    opcion_buscar_mas = opcion
                            except Exception:
                                continue

                        if opcion_buscar_mas is None:
                            for selector in Valorizado.BUSCAR_MAS:
                                try:
                                    el = find_visible_element(driver, [selector])
                                    if el is not None:
                                        opcion_buscar_mas = el
                                        break
                                except Exception:
                                    continue

                        if opcion_buscar_mas is not None:
                            self.log.info(
                                "Abriendo modal 'Buscar más...' para '%s'",
                                nombre_ubicacion,
                            )
                            try:
                                opcion_buscar_mas.click()
                            except Exception:
                                click_with_fallback(
                                    driver, opcion_buscar_mas, "Buscar mas"
                                )
                            time.sleep(1.2)
                            if self._seleccionar_ubicacion_en_modal(
                                driver, nombre_ubicacion
                            ):
                                return True
                            self.log.warning(
                                "No se selecciono '%s' desde el modal",
                                nombre_ubicacion,
                            )
                            return False

                    self.log.warning(
                        "No se encontro '%s' en el autocomplete y no hubo "
                        "opcion 'Buscar más...'",
                        nombre_ubicacion,
                    )
                    return False

                except Exception as e:
                    self.log.warning("Error buscando INPUT: %s", e)

            elif len(selects_visibles) >= 2:
                select_ubicacion = selects_visibles[1]
                select = Select(select_ubicacion)
                for opcion in select.options:
                    if opcion.text.strip() == nombre_ubicacion:
                        select.select_by_visible_text(opcion.text.strip())
                        time.sleep(0.5)
                        self.log.info(
                            "Ubicación '%s' seleccionada", nombre_ubicacion
                        )
                        return True
                self.log.warning(
                    "'%s' no encontrada en SELECT", nombre_ubicacion
                )
                return False

            self.log.error("No se pudo seleccionar la ubicación")
            return False

        except Exception as e:
            self.log.error(
                "Error después de %.1fs: %s", time.time() - tiempo_inicio, e
            )
            return False

    def _texto_opcion(self, opcion) -> str:
        import re
        texto = opcion.text.strip() if opcion.text else ""
        if not texto:
            try:
                ancla = opcion.find_element(By.XPATH, ".//a")
                texto = ancla.text.strip() if ancla.text else ""
            except Exception:
                pass
        return re.sub(r"\s+", " ", texto)

    def _seleccionar_ubicacion_en_modal(self, driver, nombre_ubicacion: str) -> bool:
        self.log.info(
            "Buscando '%s' en modal de ubicacion...", nombre_ubicacion
        )
        try:
            descartar_alert_js(driver)
            # Esperar a que el modal y su tabla esten cargados
            tabla_visible = False
            for selector in Valorizado.MODAL_ABRIR:
                try:
                    webdriver_visible = WebDriverWait(driver, 5).until(
                        EC.visibility_of_element_located(("xpath", selector))
                    )
                    if webdriver_visible is not None:
                        tabla_visible = True
                        break
                except Exception:
                    continue
            if not tabla_visible:
                time.sleep(1.5)

            filas = []
            for selector_base in Valorizado.MODAL_FILA_UBICACION:
                selector = selector_base.format(VB=nombre_ubicacion)
                try:
                    encontrados = driver.find_elements(By.XPATH, selector)
                    if encontrados:
                        filas = encontrados
                        break
                except Exception:
                    continue
            self.log.info(
                "%d fila(s) candidata(s) encontrada(s) en modal",
                len(filas),
            )
            for fila in filas:
                try:
                    if fila.is_displayed():
                        click_with_fallback(
                            driver, fila, "fila ubicacion en modal"
                        )
                        time.sleep(0.8)
                        self.log.info(
                            "Ubicación '%s' seleccionada desde modal",
                            nombre_ubicacion,
                        )
                        return True
                except Exception:
                    continue

            self.log.warning(
                "No se encontro la fila '%s' en el modal de ubicacion",
                nombre_ubicacion,
            )
            return False

        except Exception as e:
            self.log.error("Error seleccionando ubicacion en modal: %s", e)
            return False

    def seleccionar_almacen_dropdown(self, driver, nombre_almacen: str) -> bool:
        self.log.info("Seleccionando almacen '%s'...", nombre_almacen)

        tiempo_inicio = time.time()
        timeout_total = 30

        try:
            descartar_alert_js(driver)
            time.sleep(2)

            selects_en_modal = driver.find_elements(
                By.XPATH, "//div[contains(@class, 'modal')]//select"
            )
            selects_visibles = [s for s in selects_en_modal if s.is_displayed()]
            self.log.info("%d SELECT(s) visible(s)", len(selects_visibles))

            if len(selects_visibles) == 1:
                inputs = driver.find_elements(By.XPATH, Valorizado.UBICACION_INPUT)
                inputs_visibles = [inp for inp in inputs if inp.is_displayed()]
                campo_almacen = (
                    inputs_visibles[-1]
                    if len(inputs_visibles) >= 2
                    else (inputs_visibles[0] if inputs_visibles else None)
                )

                if campo_almacen:
                    driver.execute_script(
                        "arguments[0].scrollIntoView({block: 'center'});",
                        campo_almacen,
                    )
                    driver.execute_script("arguments[0].value = '';", campo_almacen)
                    time.sleep(0.3)
                    click_with_fallback(driver, campo_almacen, "campo almacen")
                    time.sleep(0.5)
                    descartar_alert_js(driver)
                    campo_almacen.send_keys(nombre_almacen)

                    opciones_encontradas = []
                    for selector in Valorizado.AUTOCOMPLETE_OPCIONES:
                        if time.time() - tiempo_inicio > timeout_total:
                            break
                        try:
                            descartar_alert_js(driver)
                            WebDriverWait(driver, 3).until(
                                EC.presence_of_element_located(("xpath", selector))
                            )
                            opciones = driver.find_elements(By.XPATH, selector)
                            opciones_visibles = [
                                opt for opt in opciones if opt.is_displayed()
                            ]
                            if opciones_visibles:
                                opciones_encontradas = opciones_visibles
                                break
                        except Exception:
                            continue

                    if opciones_encontradas:
                        for opcion in opciones_encontradas:
                            try:
                                if opcion.text.strip() == nombre_almacen:
                                    opcion.click()
                                    time.sleep(0.5)
                                    self.log.info(
                                        "Almacén '%s' seleccionado", nombre_almacen
                                    )
                                    return True
                            except Exception:
                                continue
                        try:
                            opciones_encontradas[0].click()
                            time.sleep(0.5)
                            self.log.info(
                                "Primera opción seleccionada (fallback)"
                            )
                            return True
                        except Exception:
                            pass

                    campo_almacen.send_keys(Keys.RETURN)
                    time.sleep(0.5)
                    return True

            elif len(selects_visibles) >= 2:
                select = Select(selects_visibles[1])
                for opcion in select.options:
                    if opcion.text.strip() == nombre_almacen:
                        select.select_by_visible_text(opcion.text.strip())
                        time.sleep(0.5)
                        self.log.info(
                            "Almacén '%s' seleccionado", nombre_almacen
                        )
                        return True
                return False

            return False

        except Exception as e:
            self.log.error("Error seleccionando almacen: %s", e)
            return False

    def _crear_valorizado_vacio(
        self, carpeta_destino: str, nombre_archivo: str
    ) -> str | None:
        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.cell(row=9, column=1, value="Referencia interna")
            ws.cell(row=9, column=2, value="Cantidad")
            ruta_archivo = os.path.join(carpeta_destino, nombre_archivo)
            wb.save(ruta_archivo)
            self.log.info(
                "Archivo vacío creado: %s (bodega sin unidades)", nombre_archivo
            )
            return ruta_archivo
        except Exception as e:
            self.log.error("No se pudo crear archivo vacío: %s", e)
            return None

    def generar_xlsx(self, driver) -> bool:
        self.log.info("Haciendo click en 'Generar XLSX'...")

        try:
            for selector in Valorizado.GENERAR_XLSX:
                try:
                    boton = WebDriverWait(driver, 10).until(
                        EC.element_to_be_clickable(("xpath", selector))
                    )
                    boton.click()
                    time.sleep(1)
                    self.log.info("'Generar XLSX' clickeado - Descarga iniciada")
                    return True
                except Exception:
                    continue
            raise RuntimeError("No se encontro el boton 'Generar XLSX'")
        except Exception as e:
            self.log.error("Error generando XLSX: %s", e)
            return False

    def renombrar_y_mover_archivo(
        self,
        archivo_original: str,
        nuevo_nombre: str,
        carpeta_destino: str,
    ) -> str:
        self.log.info("Renombrando archivo a: %s", nuevo_nombre)
        try:
            nueva_ruta = os.path.join(carpeta_destino, nuevo_nombre)
            if os.path.exists(nueva_ruta):
                os.remove(nueva_ruta)
            if os.path.dirname(archivo_original) != carpeta_destino:
                shutil.move(archivo_original, nueva_ruta)
            else:
                os.rename(archivo_original, nueva_ruta)
            self.log.info("Archivo guardado como: %s", nuevo_nombre)
            return nueva_ruta
        except Exception as e:
            self.log.error("Error renombrando archivo: %s", e)
            return archivo_original
